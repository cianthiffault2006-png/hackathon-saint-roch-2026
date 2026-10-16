"""Builds the final pitch deck, C:\\Users\\maeli\\cian-claude\\media\\Maillage-pitch.pptx (11 slides, 16:9), from the
teammate's structure and wording (deck-coequipiere.pdf, 16 Oct) plus the captures, the embedded tour video and
bullet-point speaker notes. Visual identity = the app's « point de maille » tokens (src/styles.css).
Run: py -3.12 docs/pitch/build_deck.py   then docs/pitch/finish_deck.ps1 (video on click, PNG renders, PDF).
"""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

MEDIA = Path(r"C:\Users\maeli\cian-claude\media")
CAP = MEDIA / "captures-d2"
IMG = Path(__file__).resolve().parent / "img"
OUT = MEDIA / "Maillage-pitch.pptx"
VIDEO = MEDIA / "tour-cadre-court.mp4"
POSTER = MEDIA / "pour-diapos-poster.png"
URL = "cianthiffault2006-png.github.io/hackathon-saint-roch-2026"

# App palette (styles.css)
PAPIER, CRAIE, LIN, TRAIT = "F7F1E6", "FFFDF8", "EFE7D8", "D9CFBD"
ENCRE, GRIS, VERT, VERT_F, VERT_P = "15201B", "4F5F58", "1F5F4A", "164536", "E6F1EC"
BRIQUE, BRIQUE_F, BRIQUE_P, OCRE = "B9460E", "8F360A", "FBE9DC", "E3B100"
TITRE_FONT, CORPS_FONT = "Georgia", "Segoe UI"   # Fraunces / Atkinson Hyperlegible are not installed on this PC
W, H = Inches(13.333), Inches(7.5)
M = Inches(0.6)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, s, size, color=ENCRE, bold=False, font=CORPS_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    lines = s if isinstance(s, list) else [s]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = rgb(color)
    return tb


def couture(slide, x, y, w):
    """Dashed green stitch line (the app's « couture »)."""
    seg, gap, cx = Inches(0.1), Inches(0.07), x
    while cx < x + w:
        rect(slide, cx, y, min(seg, x + w - cx), Inches(0.045), VERT)
        cx += seg + gap


def base(slide, n, titre=None, sous=None, dark=False):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(PAPIER)
    rect(slide, 0, 0, Inches(0.22), H, VERT)
    if titre:
        text(slide, M, Inches(0.35), W - 2 * M, Inches(0.9), titre, 38, VERT_F, True, TITRE_FONT)
        couture(slide, M + Inches(0.05), Inches(1.22), Inches(2.2))
    if sous:
        text(slide, M, Inches(1.35), W - 2 * M, Inches(0.7), sous, 20, GRIS, italic=True)
    text(slide, M, H - Inches(0.45), Inches(6), Inches(0.35), "Équipe 14 · Maillage", 11, GRIS)
    text(slide, W - M - Inches(1), H - Inches(0.45), Inches(1), Inches(0.35), str(n), 11, GRIS, align=PP_ALIGN.RIGHT)


def attribution(slide):
    text(slide, Inches(4.6), H - Inches(0.45), Inches(7.1), Inches(0.35), "Fond de carte : © OpenStreetMap, OpenFreeMap · profils fictifs et données ouvertes (CC BY 4.0)", 10, GRIS, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, titre, corps, accent=VERT, size_t=20, size_c=15, body_top=1.3):
    """body_top leaves room for a two-line title (titles wrap on the 3- and 4-card slides)."""
    rect(slide, x, y, w, h, CRAIE, TRAIT, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.06
    rect(slide, x, y, Inches(0.12), h, accent)
    text(slide, x + Inches(0.22), y + Inches(0.15), w - Inches(0.4), Inches(1.1), titre, size_t, VERT_F, True, TITRE_FONT)
    text(slide, x + Inches(0.22), y + Inches(body_top), w - Inches(0.4), h - Inches(body_top + 0.15), corps, size_c, ENCRE)


def picture_fit(slide, path, x, y, max_w, max_h):
    """Adds a picture scaled to fit the box, centred in it."""
    pic = slide.shapes.add_picture(str(path), x, y)
    ratio = min(max_w / pic.width, max_h / pic.height)
    pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
    pic.left = int(x + (max_w - pic.width) / 2)
    pic.top = int(y + (max_h - pic.height) / 2)
    return pic


def notes(slide, who, bullets, cue):
    tf = slide.notes_slide.notes_text_frame
    tf.text = f"{who}"
    for b in bullets:
        tf.add_paragraph().text = f"• {b}"
    tf.add_paragraph().text = f"→ {cue}"


# 1. Titre
s = prs.slides.add_slide(BLANK)
base(s, 1)
text(s, M + Inches(0.3), Inches(1.6), Inches(7.2), Inches(1.5), "Maillage", 80, VERT_F, True, TITRE_FONT)
couture(s, M + Inches(0.4), Inches(3.05), Inches(2.4))
text(s, M + Inches(0.3), Inches(3.3), Inches(7.2), Inches(1.4), "Un café, une artiste, un organisme : un projet.", 30, ENCRE, font=TITRE_FONT)
text(s, M + Inches(0.3), Inches(5.1), Inches(7.2), Inches(0.6), "Équipe 14 · Hackathon étudiant 2026 · Saint-Roch", 16, GRIS)
picture_fit(s, CAP / "phone-accueil.png", Inches(8.6), Inches(0.5), Inches(4.2), Inches(6.4))
attribution(s)
notes(s, "COÉQUIPIÈRE (8 s)", ["Maillage, équipe 14", "une phrase : un café, une artiste, un organisme, un projet"], "diapo 2 sur « projet »")

# 2. Le problème
s = prs.slides.add_slide(BLANK)
base(s, 2, "Le problème", "Les initiatives existent. Ce qui manque, c’est le réseau qui les relie.")
cards = [
    ("Espaces inexploités", "Commerçants disposant de locaux libres sans connaître les artistes locaux.", BRIQUE),
    ("Organismes isolés", "Projets communautaires en manque de visibilité et de partenaires adaptés.", "0072B2"),
    ("Citoyens engagés", "Volonté d’agir pour le quartier sans canal direct ni structure d’organisation.", "7B3F9E"),
]
cw = (W - 2 * M - 2 * Inches(0.35)) / 3
for i, (t, c, a) in enumerate(cards):
    card(s, M + i * (cw + Inches(0.35)), Inches(2.35), cw, Inches(3.3), t, c, a, 22, 17)
text(s, M, Inches(5.95), W - 2 * M, Inches(0.6), "Le défi du cahier : connecter la communauté de Saint-Roch, des acteurs qui manquent d’outils pour collaborer.", 15, GRIS, italic=True)
notes(s, "COÉQUIPIÈRE (22 s)", ["le défi du cahier : connecter la communauté, des acteurs sans outils pour collaborer", "ressources là : locaux, talent, bénévoles; manque : savoir qui a quoi, juste à côté", "Mathieu (café, jeudi soir libre), Camille (cherche un mur), le Collectif (cherche un public) : voisins, ne se connaissent pas", "jamais : « on nous a dit », aucun vrai nom"], "diapo 3 sur « existent »")

# 3. La solution
s = prs.slides.add_slide(BLANK)
base(s, 3, "La solution : Maillage", "Un besoin en mots simples → trois partenaires complémentaires, les raisons, la distance à pied.")
picture_fit(s, CAP / "phone-proposition-score.png", M, Inches(2.1), Inches(2.6), Inches(4.7))
picture_fit(s, CAP / "desktop-proposition-score.png", Inches(3.5), Inches(2.1), Inches(9.2), Inches(4.7))
rect(s, Inches(9.35), Inches(6.35), Inches(3.35), Inches(0.5), VERT_P, VERT, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.5
text(s, Inches(9.35), Inches(6.35), Inches(3.35), Inches(0.5), "Testé sur téléphone et ordinateur", 14, VERT_F, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
attribution(s)
notes(s, "COÉQUIPIÈRE (22 s)", ["l’app qui propose à ces trois-là un projet commun", "on écrit son besoin en mots simples → trois partenaires, à distance de marche, avec les raisons; chacun accepte ou refuse", "ne remplace aucun outil : crée le lien qui manque", "testée sur téléphone et ordinateur; quatre étudiants, équipe 14"], "diapo 4 sur « équipe 14 » (Cian)")

# 4. Moteur de recommandation
s = prs.slides.add_slide(BLANK)
base(s, 4, "Moteur de recommandation", "Un score pondéré, visible dans l’app : chaque suggestion s’explique.")
picture_fit(s, CAP / "phone-proposition-score.png", M, Inches(2.1), Inches(2.7), Inches(4.8))
poids = [("Complémentarité des besoins", 35), ("Réciprocité des contributions", 20), ("Proximité à pied", 20), ("Thème commun", 15), ("Disponibilités compatibles", 10)]
bx, by, bw = Inches(3.8), Inches(2.25), Inches(6.3)
for i, (lab, p) in enumerate(poids):
    y = by + i * Inches(0.8)
    text(s, bx, y, Inches(3.3), Inches(0.5), lab, 16, ENCRE, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, bx + Inches(3.4), y + Inches(0.08), bw - Inches(3.4), Inches(0.34), LIN, TRAIT)
    rect(s, bx + Inches(3.4), y + Inches(0.08), int((bw - Inches(3.4)) * p / 35), Inches(0.34), VERT if i else BRIQUE)
    text(s, bx + bw + Inches(0.1), y, Inches(1.2), Inches(0.5), f"{p} %", 18, VERT_F, True, anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(3.8), Inches(6.35), Inches(9), Inches(0.6), "Dans l’app : « Comment ce score », « Pourquoi cette suggestion », « À confirmer ». Règles lisibles, pas d’IA payante.", 14, GRIS, italic=True)
attribution(s)
notes(s, "CIAN (18 s)", ["cinq critères, poids visibles : 35 / 20 / 20 / 15 / 10, nos choix, affichés pour être débattus", "exemple : Réciprocité 58 % = le café comble une partie des besoins de ses partenaires, pas tout, et on le dit", "ce qu’on ne sait pas, « À confirmer »", "IA testée (/labo/) : profils semblables, pas complémentaires → règles aujourd’hui, IA en phase 2"], "diapo 5 sur « phase 2 »")

# 5. Technologies, outils et données
s = prs.slides.add_slide(BLANK)
base(s, 5, "Technologies, outils et données utilisés", "Personne ne commence avec une page blanche.")
cards5 = [
    ("Technologies", "React, Vite, TypeScript, MapLibre GL, GitHub Pages. Claude Code comme assistant de code.", VERT),
    ("Données ouvertes", "Ville de Québec et SIT Québec, via Données Québec, licence CC BY 4.0 : lieux publics, événements, vrais commerces.", "00694D"),
    ("Répertoires fournis", "386 membres (fichier fourni par la SDC), 614 organismes reconnus (toute la ville) : colonnes et totaux seulement, aucune valeur en ligne.", BRIQUE),
    ("Loi 25", "Fiche publique seulement après réclamation, code à l’adresse déjà au dossier et consentement. Non réclamée = visible par la SDC seulement (pilote).", "7B3F9E"),
]
cw = (W - 2 * M - 3 * Inches(0.25)) / 4
for i, (t, c, a) in enumerate(cards5):
    card(s, M + i * (cw + Inches(0.25)), Inches(2.2), cw, Inches(4.2), t, c, a, 19, 14)
notes(s, "CIAN (30 s)", ["données ouvertes Ville + SIT Québec, CC BY", "répertoires qui existent déjà : 386 membres (fichier SDC), 614 organismes reconnus à l’échelle de la ville (jamais « dans Saint-Roch ») : colonnes seulement, aucune valeur confidentielle en ligne", "Loi 25 : le consentement décide; public après réclamation, code à l’adresse au dossier, consentement; sinon SDC seulement en pilote", "dans la démo : fiches fictives + commerces des données ouvertes, étiquetées « Pré-rempli · visible ici pour la démo »", "Claude Code a aidé à coder; les choix sont les nôtres", "[CONDITIONNEL] mode présentation local, seulement si l’organisateur a dit oui"], "diapo 6 sur « nôtres » : cadrage puis clic sur la vidéo")

# 6. Démonstration (vidéo)
s = prs.slides.add_slide(BLANK)
base(s, 6)
text(s, M, Inches(0.25), W - 2 * M, Inches(0.6), "Démonstration du prototype · enregistrée ce matin sur l’app en ligne, visite guidée intégrée", 16, GRIS, align=PP_ALIGN.CENTER)
vh = H - Inches(1.3)
vw = int(vh * 16 / 9)
vx = int((W - vw) / 2)
if VIDEO.exists():
    s.shapes.add_movie(str(VIDEO), vx, Inches(0.85), vw, vh, poster_frame_image=str(POSTER) if POSTER.exists() else None, mime_type="video/mp4")
else:
    rect(s, vx, Inches(0.85), vw, vh, LIN, TRAIT)
    text(s, vx, Inches(3), vw, Inches(1), "VIDÉO MANQUANTE : tour-cadre-court.mp4", 28, BRIQUE, True, align=PP_ALIGN.CENTER)
attribution(s)
notes(s, "CIAN puis COÉQUIPIÈRE puis CIAN (8 s + vidéo 1 min 32 s)", ["avant le clic : « Enregistré ce matin sur l’app en ligne; les clics sont ceux de la visite guidée intégrée, le lien est à l’écran »", "Cian = Mathieu (0:00 café / 0:05 j’écris / 0:10 Maillage comprend / 0:19 un projet, pas une liste / 0:25 Réciprocité 58 % / 0:31 chacun y gagne / 0:39 ce qu’on ne sait pas / 0:46 Ça m’intéresse)", "Coéquipière = Camille (0:53 je reçois, 22 oct 19 h à 22 h / 1:00 je sais pourquoi / 1:06 J’accepte / 1:10 Collectif aussi, confirmé; aucun message ne part, en pilote l’invitation part vraiment)", "Cian (1:15 la fiche / 1:30 un café, une artiste, un organisme : un projet)", "on ne lit pas les sous-titres; silence entre les repères", "plan B : Alt+Tab, Ctrl+1, « Visite guidée »; plan C : Ctrl+3 vidéo de secours"], "diapo 7 sur la carte de fin")

# 7. Impact
s = prs.slides.add_slide(BLANK)
base(s, 7, "Impact et feuille de route", "Un pilote proposé avec une SDC, puis tout Saint-Roch.")
cards7 = [
    ("Vitalité commerciale", "Génération d’activités inédites attirant des visiteurs et dynamisant les commerces locaux.", BRIQUE),
    ("Cohésion sociale", "Création de relations durables et renforcement du sentiment d’appartenance au quartier.", "0072B2"),
    ("Valorisation des actifs", "Optimisation des espaces sous-utilisés, des compétences et du matériel disponible.", VERT),
]
cw = (W - 2 * M - 2 * Inches(0.35)) / 3
for i, (t, c, a) in enumerate(cards7):
    card(s, M + i * (cw + Inches(0.35)), Inches(2.2), cw, Inches(3.3), t, c, a, 22, 17)
text(s, M, Inches(5.8), W - 2 * M, Inches(0.6), "Mesures du pilote : projets confirmés · nouveaux partenariats · fréquentation des événements créés", 16, VERT_F, True)
notes(s, "CIAN (18 s)", ["plus de projets communs, des commerces plus vivants, contribution possible au sentiment de sécurité", "on le mesure : projets confirmés, nouveaux partenariats, fréquentation", "aucun chiffre non prouvé; jamais « on a rencontré la SDC »"], "diapo 8 sur « fréquentation »")

# 8. Faisabilité et évolutions
s = prs.slides.add_slide(BLANK)
base(s, 8, "Faisabilité et évolutions futures", "Trois étapes, un pilote d’un trimestre pour commencer.")
etapes = [
    ("Étape 1 · Projet pilote SDC", "Ce qu’on propose : 10 membres volontaires d’une SDC, un trimestre, une personne à temps partiel, un export des membres. Profils activés par leurs propriétaires, organisations validées."),
    ("Étape 2 · Coordination réelle", "Invitations effectives, fil de projet modéré par la SDC (tableau de modération prévu, pas construit), notifications."),
    ("Étape 3 · Déploiement", "Modèle de langage en phase 2 (toutes langues, toutes fautes); extension éventuelle à d’autres quartiers de la ville de Québec."),
]
for i, (t, c) in enumerate(etapes):
    y = Inches(2.2) + i * Inches(1.45)
    rect(s, M, y, Inches(0.9), Inches(0.9), BRIQUE if i == 0 else VERT, shape=MSO_SHAPE.OVAL)
    text(s, M, y, Inches(0.9), Inches(0.9), str(i + 1), 28, CRAIE, True, TITRE_FONT, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, M + Inches(1.15), y - Inches(0.05), W - 2 * M - Inches(1.2), Inches(0.5), t, 21, VERT_F, True, TITRE_FONT)
    text(s, M + Inches(1.15), y + Inches(0.45), W - 2 * M - Inches(1.2), Inches(0.9), c, 15, ENCRE)
notes(s, "CIAN (18 s)", ["voici le pilote qu’on propose : dix membres volontaires d’une SDC, un trimestre, une personne à temps partiel, un export du fichier des membres, trois chiffres pour décider de la suite", "ensuite : modération (prévue, pas construite), IA en phase 2, autres quartiers", "une proposition, pas une entente"], "diapo 9 sur « quartiers » (ou diapo 10 si Q montre 2 doigts)")

# 9. Vous vous demandez peut-être…
s = prs.slides.add_slide(BLANK)
base(s, 9, "Vous vous demandez peut-être…")
qa = [
    ("Et la Loi 25?", "Public seulement après réclamation, code à l’adresse déjà au dossier et consentement; sinon, en pilote, la SDC seulement."),
    ("Où est l’IA?", "Testée : elle trouve des profils semblables, pas complémentaires. Des règles qui s’expliquent aujourd’hui; l’IA en phase 2."),
    ("Ces restaurants ont-ils accepté?", "Données ouvertes du ministère du Tourisme, affichées pré-remplies pour la démo : exactement ce que la réclamation règle en pilote."),
]
for i, (q, a) in enumerate(qa):
    y = Inches(1.7) + i * Inches(1.6)
    rect(s, M, y, W - 2 * M, Inches(1.35), CRAIE, TRAIT, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.08
    text(s, M + Inches(0.25), y + Inches(0.1), Inches(4.2), Inches(1.1), q, 22, BRIQUE_F, True, TITRE_FONT, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M + Inches(4.5), y + Inches(0.1), W - 2 * M - Inches(4.8), Inches(1.15), a, 16, ENCRE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "COÉQUIPIÈRE demande, CIAN répond (28 s) · cette diapo saute si on dépasse 4:45", ["Loi 25 → public après réclamation, code, consentement; sinon SDC seulement en pilote", "où est l’IA → testée, profils semblables pas complémentaires; règles qui s’expliquent; IA phase 2", "restaurants → données ouvertes du ministère du Tourisme, pré-remplies pour la démo; la réclamation règle ça en pilote"], "diapo 10 sur « pilote »")

# 10. Citation de fin
s = prs.slides.add_slide(BLANK)
base(s, 10)
rect(s, M + Inches(0.4), Inches(1.6), Inches(0.12), Inches(3.6), BRIQUE)
text(s, M + Inches(0.8), Inches(1.5), W - 2 * M - Inches(1.2), Inches(3.9), "« Saint-Roch possède déjà les ressources, les talents et les idées. Maillage crée les connexions nécessaires pour les transformer en projets collectifs. »", 34, VERT_F, font=TITRE_FONT, anchor=MSO_ANCHOR.MIDDLE)
text(s, M + Inches(0.8), Inches(5.6), Inches(8), Inches(0.5), "Équipe 14", 18, GRIS)
notes(s, "COÉQUIPIÈRE (8 s)", ["la citation, lue calmement", "« Le lien est à l’écran, essayez-le. Merci! »"], "diapo 11 sur « Merci! », puis silence")

# 11. Questions? QR
s = prs.slides.add_slide(BLANK)
base(s, 11)
text(s, M, Inches(0.5), Inches(6.5), Inches(1.2), "Questions?", 54, VERT_F, True, TITRE_FONT)
text(s, M, Inches(1.8), Inches(6.5), Inches(1.0), "Essayez Maillage maintenant :", 22, ENCRE)
text(s, M, Inches(2.6), Inches(6.5), Inches(1.2), URL, 20, BRIQUE_F, True)
text(s, M, Inches(4.0), Inches(6.5), Inches(1.6), ["Un café, une artiste, un organisme : un projet.", "Équipe 14 · Hackathon étudiant 2026 · Saint-Roch"], 16, GRIS)
qr = IMG / "slide-7.png"
if qr.exists():
    rect(s, Inches(7.6), Inches(0.75), Inches(5.3), Inches(5.3), CRAIE, TRAIT, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.04
    pic = s.shapes.add_picture(str(qr), Inches(7.85), Inches(1.0), Inches(4.8), Inches(4.8))
text(s, Inches(7.6), Inches(6.15), Inches(5.3), Inches(0.4), "Scannez pour ouvrir l’app (GitHub Pages)", 13, GRIS, align=PP_ALIGN.CENTER)
text(s, M, H - Inches(0.85), W - 2 * M - Inches(1.2), Inches(0.4), "Outils d’IA : Claude Code et Claude (code, scripts de données, brouillons de textes). Fait par l’équipe : cadrage, données, conception, tests, présentation.", 10, GRIS)
notes(s, "TOUS (questions, 3 min)", ["la diapo reste pendant les questions (QR); pour montrer un écran : Alt+Tab, Ctrl+1 (app, disposition bureau), puis Alt+Tab pour revenir", "une voix par réponse, 20 s; « nous ne savons pas encore » plutôt qu’inventer", "jamais la liste de réclamation à l’écran (vrais noms)"], "fin")

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"wrote {OUT} ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
