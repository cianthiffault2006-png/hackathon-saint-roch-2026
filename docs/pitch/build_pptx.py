"""Regenerate docs/pitch/Maillage.pptx from docs/pitch/slides-fr.md.

Run:  py -3.12 docs/pitch/build_pptx.py
Needs: py -3.12 -m pip install --user python-pptx

Per slide in slides-fr.md ("## <n>. <titre>"):
  Ligne : ...   main message (shown big)
  Diapo : a ; b   short points (2 to 5 become cards)
  every other line (Visuel, Signal, details) goes to the speaker notes.
Screenshots: docs/pitch/img/slide-<n>.png is placed on slide <n>.
"""
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
SRC = HERE / "slides-fr.md"
OUT = HERE / "Maillage.pptx"
IMG = HERE / "img"

GREEN = RGBColor(0x1F, 0x5F, 0x4A)
ORANGE = RGBColor(0xE8, 0x83, 0x3A)
BG = RGBColor(0xFA, 0xF7, 0xF0)
INK = RGBColor(0x22, 0x2B, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.7)


def parse(path):
    slides, cur = [], None
    for raw in path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^## (\d+)\.\s*(.+)$", raw)
        if m:
            cur = {"n": int(m.group(1)), "titre": m.group(2).strip(), "ligne": "", "points": [], "notes": []}
            slides.append(cur)
            continue
        if raw.startswith("## "):  # non numbered heading ends the slides
            cur = None
            continue
        if cur is None or not raw.strip():
            continue
        if raw.startswith("Ligne :"):
            cur["ligne"] = raw.split(":", 1)[1].strip().strip("«» ").strip()
        elif raw.startswith("Diapo :"):
            cur["points"] = [p.strip() for p in raw.split(":", 1)[1].split(";") if p.strip()]
        else:
            cur["notes"].append(raw.strip())
    return slides


def set_bg(slide):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = BG


def text(slide, x, y, w, h, s, size, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = FONT
    r.font.color.rgb = color
    return tb


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def place_image(slide, path, x, y, w, h):
    from pptx.parts.image import Image

    iw, ih = Image.from_file(str(path)).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(str(path), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)


def footer(slide, n, with_attrib):
    text(slide, MARGIN, H - Inches(0.5), Inches(8), Inches(0.35),
         "Équipe 14 · Maillage" + ("  ·  Fond de carte : © OpenStreetMap, OpenFreeMap" if with_attrib else ""),
         12, GREEN)
    text(slide, W - MARGIN - Inches(1), H - Inches(0.5), Inches(1), Inches(0.35), str(n), 12, GREEN,
         align=PP_ALIGN.RIGHT)


def build_title(prs, s, img):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    rect(slide, 0, 0, Inches(0.35), H, ORANGE)
    left_w = Inches(7.4) if img else W - 2 * MARGIN
    text(slide, MARGIN, Inches(1.9), left_w, Inches(1.6), "Maillage", 88, GREEN, bold=True)
    rect(slide, MARGIN + Inches(0.05), Inches(3.6), Inches(1.6), Inches(0.08), ORANGE)
    text(slide, MARGIN, Inches(3.85), left_w, Inches(1.4), s["ligne"], 32, INK)
    text(slide, MARGIN, Inches(5.6), left_w, Inches(0.5),
         "Équipe 14 · Hackathon étudiant 2026 · Saint-Roch", 20, GREEN)
    if img:
        place_image(slide, img, Inches(8.4), Inches(0.8), Inches(4.4), Inches(5.9))
    return slide


def build_content(prs, s, img):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    rect(slide, 0, 0, W, Inches(1.45), GREEN)
    text(slide, MARGIN, Inches(0.2), W - 2 * MARGIN, Inches(1.05), s["titre"], 40, WHITE, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    rect(slide, 0, Inches(1.45), W, Inches(0.07), ORANGE)

    top = Inches(1.95)
    body_h = H - top - Inches(0.8)
    text_w = Inches(6.9) if img else W - 2 * MARGIN
    n_pts = len(s["points"])

    if s["ligne"]:
        text(slide, MARGIN, top, text_w, Inches(1.5), s["ligne"], 30 if img else 34, INK, bold=True)
    pts_top = top + Inches(1.7)

    if 2 <= n_pts <= 5:
        if img:  # vertical stack beside the image
            ch = min(Inches(1.0), (H - pts_top - Inches(0.8)) // n_pts - Inches(0.12))
            for i, p in enumerate(s["points"]):
                y = pts_top + i * (ch + Inches(0.12))
                rect(slide, MARGIN, y, text_w, ch, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE)
                rect(slide, MARGIN, y, Inches(0.12), ch, ORANGE)
                text(slide, MARGIN + Inches(0.3), y, text_w - Inches(0.4), ch, p, 20, INK,
                     anchor=MSO_ANCHOR.MIDDLE)
        else:  # cards in a row
            gap = Inches(0.25)
            cw = (W - 2 * MARGIN - gap * (n_pts - 1)) // n_pts
            ch = Inches(2.6)
            size = 24 if n_pts <= 3 else 20
            for i, p in enumerate(s["points"]):
                x = MARGIN + i * (cw + gap)
                rect(slide, x, pts_top, cw, ch, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE)
                rect(slide, x, pts_top, cw, Inches(0.12), ORANGE)
                text(slide, x + Inches(0.2), pts_top + Inches(0.25), cw - Inches(0.4), ch - Inches(0.4),
                     p, size, INK, anchor=MSO_ANCHOR.MIDDLE)
    elif n_pts == 1:
        text(slide, MARGIN, pts_top, text_w, Inches(1.5), s["points"][0], 24, GREEN)

    if img:
        place_image(slide, img, Inches(8.0), top - Inches(0.1), Inches(4.7), body_h + Inches(0.1))
    footer(slide, s["n"], bool(img))
    return slide


def main():
    slides = parse(SRC)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    for s in slides:
        img = IMG / f"slide-{s['n']}.png"
        img = img if img.exists() else None
        slide = build_title(prs, s, img) if s["n"] == 1 else build_content(prs, s, img)
        slide.notes_slide.notes_text_frame.text = "\n".join(s["notes"])
        words = len(" ".join([s["ligne"]] + s["points"]).split())
        if words > 30:
            print(f"warning: slide {s['n']} has {words} words on the slide")
    prs.save(OUT)
    print(f"wrote {OUT} ({len(slides)} slides)")


if __name__ == "__main__":
    main()
