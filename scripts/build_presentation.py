"""Mode présentation (local only): turns the organizers' provided directories into app rows for Cian's laptop.
Writes <out>/profils-reels.json, lieux-reels.json, evenements-reels.json (git-ignored, served by Vite in dev/preview
only, never built into dist/). Prints counts and column names only, never a cell value.
Usage: py -3.12 scripts/build_presentation.py [--source data-private] [--out data-private/app] [--adresses <csv>]
  --adresses: the open « Adresses de la Ville de Québec » CSV (CC BY 4.0); downloaded on first run if missing.
Business and organisation fields only: no person name, phone or email is ever read into the output.
"""
import argparse, csv, hashlib, json, re, sys, unicodedata, urllib.request
from pathlib import Path
from shapely.geometry import shape, Point

ROOT = Path(__file__).resolve().parent.parent
ADRESSES_URL = ("https://www.donneesquebec.ca/recherche/dataset/2567eac8-a3ce-423a-a43d-6b8e74708c1d/resource/"
                "cef2c0b7-8eac-4466-9a0b-6987501edace/download/vdq-adresse.csv")
EI, EE, V, A, P, CA, AN = "espace_interieur", "espace_exterieur", "visibilite", "affichage", "public", "contenu_artistique", "animation"
FOURNI = "Données fournies par la Ville de Québec (entente de confidentialité) : affichage local seulement"

# file -> sheet column names used (names only; from scripts/schema.py)
FICHIERS = {
    "entreprises": ("DEGP - Entreprises St-Roch-Rev.xlsx", ["Type rue", "Nom rue", "Marque de Commerce"]),
    "sdc": ("DEGP - Membres - SDC Centre-Ville-Rev.xlsx", ["Type rue", "Nom rue", "OBNL", "Marque de Commerce", "Activité commerciale"]),
    "reconnus": ("CP-LSVC-Liste organismes reconnus LSVC-k.xlsx", ["Nom", "Domaine d'intervention", "Territoire(s)", "Statut", "Catégorie"]),
    "installations": ("CP-LSVC-Liste installations-k.xlsx", ["Nom_Type_No_Installation_calculé", "Adresse", "Type d'équipement texte"]),
    "art": ("CP-LSVC-Collection Art Public-k.xlsx", ["Titre œuvre", "Quartier ", "Adresse civique", "Géocode", "Type emplacement"]),
    "evenements": ("CP-LSVC-Événements à venir-k.xlsx", ["Titre PROJET", "Catégorie", "Site(s) de l'activité", "Quartier", "Date et heure de début", "Date et heure de fin"]),
}
TYPES_VOIE = {"rue", "boulevard", "boul", "avenue", "av", "chemin", "ch", "cote", "côte", "place", "pl", "ruelle", "allee", "allée", "promenade", "carre", "carré"}


def norm(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9 ]+", " ", s).strip()


ARTICLES = {"de", "la", "du", "des", "le", "les", "d", "l"}
CANON = {"est": "e", "ouest": "o", "st": "saint", "ste": "sainte", "nord": "n", "sud": "s"}


def mots_rue(s) -> tuple:
    """Street name as a canonical tuple: no type words, no articles, 'Est' -> 'e', 'St' -> 'saint'."""
    return tuple(CANON.get(w, w) for w in norm(s).split() if w not in TYPES_VOIE and w not in ARTICLES)


def hid(*parts) -> str:
    return hashlib.sha1("|".join(str(p) for p in parts).encode()).hexdigest()[:10]


class Adresses:
    """Open address points inside the Saint-Roch outline: street -> points, (civic, street) -> point."""

    def __init__(self, csv_path: Path, poly):
        self.par_rue, self.par_civique, self.tous = {}, {}, []
        inner = poly.buffer(-0.0003)
        for r in csv.DictReader(open(csv_path, encoding="utf-8-sig")):
            try:
                x, y = float(r["LONGITUDE"]), float(r["LATITUDE"])
            except ValueError:
                continue
            if not poly.contains(Point(x, y)):
                continue
            # NOM_VOIE already holds the link words ("DE LA COURONNE"); orientation is E/O
            rue = mots_rue(f"{r['NOM_VOIE']} {r.get('ORIENTATION', '')}")
            rue_sans = mots_rue(r["NOM_VOIE"])
            pt = (round(y, 6), round(x, 6))
            for k in {rue, rue_sans}:
                self.par_rue.setdefault(k, []).append(pt)
                self.par_civique[(norm(r["NO_CIVIQUE"]), k)] = pt
            if inner.contains(Point(x, y)):
                self.tous.append(pt)
        for k in self.par_rue:
            self.par_rue[k].sort()

    def rue(self, type_rue, nom_rue):
        for k in (mots_rue(f"{type_rue} {nom_rue}"), mots_rue(nom_rue)):
            if k in self.par_rue:
                return self.par_rue[k]
        return None

    def adresse(self, texte):
        """'123, rue Untel Est' -> point, or None."""
        m = re.match(r"\s*(\d+)[a-zA-Z]?\s*[,\-]?\s*(.+)", str(texte or ""))
        if not m:
            return None
        civ, reste = m.group(1), m.group(2).split(",")[0]
        k = mots_rue(reste)
        pt = self.par_civique.get((civ, k))
        if pt:
            return pt
        # closest civic number on the same street
        cands = [(abs(int(c) - int(civ)), p) for (c, kk), p in self.par_civique.items() if kk == k and c.isdigit()]
        return min(cands)[1] if cands else None


def lire_xlsx(path: Path, colonnes):
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    entetes = [str(h).strip() if h is not None else "" for h in next(rows)]
    idx = {}
    for c in colonnes:
        c_norm = c.strip()
        trouve = [i for i, h in enumerate(entetes) if h == c_norm]
        if not trouve:
            raise KeyError(f"colonne absente: {c!r} (colonnes du fichier: {entetes})")
        idx[c_norm] = trouve[0]
    out = []
    for r in rows:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({c: (r[i] if i < len(r) else None) for c, i in idx.items()})
    return out


def texte(v) -> str:
    return "" if v is None else str(v).strip()


def spread(n, pts, seed):
    """n deterministic points spread over the outline (indicative positions for rows without an address)."""
    if not pts or n == 0:
        return []
    start = int(hid(seed), 16) % len(pts)
    chosen = [pts[start]]
    while len(chosen) < n:
        chosen.append(max(pts, key=lambda p: min((p[0] - c[0]) ** 2 + (p[1] - c[1]) ** 2 for c in chosen)))
    return chosen


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", default=str(ROOT / "data-private"))
    ap.add_argument("--out", default=str(ROOT / "data-private" / "app"))
    ap.add_argument("--adresses", default=str(ROOT / "data-private" / "adresses" / "vdq-adresse.csv"))
    a = ap.parse_args()
    src, out, adr = Path(a.source), Path(a.out), Path(a.adresses)
    if not adr.exists():
        adr.parent.mkdir(parents=True, exist_ok=True)
        print("téléchargement des adresses ouvertes (Ville de Québec, CC BY 4.0)...")
        urllib.request.urlretrieve(ADRESSES_URL, adr)
    poly = shape(json.load(open(ROOT / "data/open/saint_roch_quartier.geojson", encoding="utf-8"))["features"][0]["geometry"])
    A_ = Adresses(adr, poly)
    print(f"adresses ouvertes dans Saint-Roch: {len(A_.tous)} points, {len(A_.par_rue)} clés de rue")
    profils, lieux, evs = [], [], []
    rapport = []

    def fichier(cle):
        nom, cols = FICHIERS[cle]
        p = src / nom
        if not p.exists():
            rapport.append(f"{nom}: absent, ignoré")
            return None
        rows = lire_xlsx(p, cols)
        rapport.append(f"{nom}: {len(rows)} lignes lues, colonnes {cols}")
        return rows

    # 1-2. commerces: street only (no civic number in the file) -> spread along the street inside the outline
    for cle, origine, verifie in (("entreprises", "Entreprises St-Roch", "non"), ("sdc", "Membres SDC Centre-Ville", "sdc")):
        rows = fichier(cle)
        if rows is None:
            continue
        drop = {"nom vide ou non textuel": 0, "rue hors Saint-Roch ou introuvable": 0, "doublon": 0}
        vus = set()
        par_rue_compteur = {}
        for r in rows:
            nom = texte(r["Marque de Commerce"])
            if not nom or not isinstance(r["Marque de Commerce"], str):
                drop["nom vide ou non textuel"] += 1
                continue
            pts = A_.rue(r["Type rue"], r["Nom rue"])
            if not pts:
                drop["rue hors Saint-Roch ou introuvable"] += 1
                continue
            k = norm(nom)
            if k in vus:
                drop["doublon"] += 1
                continue
            vus.add(k)
            i = par_rue_compteur.get(id(pts), 0)
            par_rue_compteur[id(pts)] = i + 1
            pt = pts[(i * 7) % len(pts)]
            activite = texte(r.get("Activité commerciale")) if cle == "sdc" else ""
            desc = (f"Membre de la SDC Centre-Ville. Activité : {activite}. " if activite else "Commerce de Saint-Roch. ") + \
                "Position approximative sur la rue (le répertoire n'a pas de numéro civique). " + FOURNI + "."
            profils.append(dict(id=f"fourni-{hid(cle, nom)}", role="commerce", sousType=(activite.lower()[:40] or "commerce"), nom=nom,
                                description=desc, lat=pt[0], lng=pt[1], offres=[V], besoins=[AN, V], creneaux=["semaine_jour", "semaine_soir"],
                                themes=["commerce_local"], source="fourni", origine=origine, reclame=False, verifie=verifie, contactPublic=False))
        rapport.append(f"  -> {len(vus)} profils écrits; rejets: {drop}")

    # 3. organismes reconnus: no address -> indicative positions; only those whose territory names the quartier
    rows = fichier("reconnus")
    if rows is not None:
        gard, drop = [], {"nom vide": 0, "territoire sans Saint-Roch": 0, "doublon": 0}
        vus = set()
        for r in rows:
            nom = texte(r["Nom"])
            if not nom:
                drop["nom vide"] += 1
                continue
            terr = norm(r["Territoire(s)"])
            if "saint roch" not in terr and "st roch" not in terr:
                drop["territoire sans Saint-Roch"] += 1
                continue
            if norm(nom) in vus:
                drop["doublon"] += 1
                continue
            vus.add(norm(nom))
            gard.append(r)
        for r, pt in zip(gard, spread(len(gard), A_.tous, "reconnus")):
            dom, cat = texte(r["Domaine d'intervention"]), texte(r["Catégorie"])
            desc = f"Organisme reconnu par la Ville. Domaine : {dom or 'non précisé'}. Catégorie : {cat or 'non précisée'}. " \
                   "Adresse non fournie : position indicative dans le quartier. " + FOURNI + "."
            profils.append(dict(id=f"fourni-{hid('reconnus', r['Nom'])}", role="organisme", sousType=(dom.lower()[:40] or "organisme"), nom=texte(r["Nom"]),
                                description=desc, lat=pt[0], lng=pt[1], offres=[P, "benevoles"], besoins=[V, EI], creneaux=["semaine_jour"],
                                themes=["communautaire"], source="fourni", origine="Liste organismes reconnus LSVC", reclame=False, verifie="ville", contactPublic=False))
        rapport.append(f"  -> {len(gard)} profils écrits; rejets: {drop}")

    # 4. installations: geocoded by address, inside the outline
    rows = fichier("installations")
    if rows is not None:
        drop = {"adresse introuvable ou hors Saint-Roch": 0, "nom vide": 0}
        n = 0
        for r in rows:
            nom = texte(r["Nom_Type_No_Installation_calculé"])
            if not nom:
                drop["nom vide"] += 1
                continue
            pt = A_.adresse(r["Adresse"])
            if not pt:
                drop["adresse introuvable ou hors Saint-Roch"] += 1
                continue
            t = norm(r["Type d'équipement texte"])
            typ, of = ("parc", [EE, P]) if re.search(r"parc|jardin|terrain|piste|patinoire", t) else ("stationnement", []) if "stationnement" in t else ("salle", [EI, P])
            lieux.append(dict(id=f"fourni-{hid('inst', nom)}", nom=nom, type=typ, lat=pt[0], lng=pt[1], offres=of, source="fourni", attribution=FOURNI))
            n += 1
        rapport.append(f"  -> {n} lieux écrits; rejets: {drop}")

    # 5. art public: title only (no artist name), Saint-Roch rows, position from Géocode or the civic address
    rows = fichier("art")
    if rows is not None:
        drop = {"quartier autre": 0, "titre vide": 0, "position introuvable": 0}
        n = 0
        for r in rows:
            if "saint roch" not in norm(r["Quartier"]) and "st roch" not in norm(r["Quartier"]):
                drop["quartier autre"] += 1
                continue
            titre = texte(r["Titre œuvre"])
            if not titre:
                drop["titre vide"] += 1
                continue
            pt = None
            nums = re.findall(r"-?\d+\.\d+", texte(r["Géocode"]))
            if len(nums) >= 2:
                f1, f2 = float(nums[0]), float(nums[1])
                lat, lng = (f1, f2) if 40 < f1 < 50 else (f2, f1)
                if poly.contains(Point(lng, lat)):
                    pt = (round(lat, 6), round(lng, 6))
            pt = pt or A_.adresse(r["Adresse civique"])
            if not pt:
                drop["position introuvable"] += 1
                continue
            lieux.append(dict(id=f"fourni-{hid('art', titre, r['Adresse civique'])}", nom=f"Œuvre d'art public : {titre}", type="lieu_culturel",
                              lat=pt[0], lng=pt[1], offres=[V], source="fourni", attribution=FOURNI))
            n += 1
        rapport.append(f"  -> {n} lieux écrits; rejets: {drop}")

    # 6. upcoming events in Saint-Roch: title, category, site, dates (promoter and descriptions not read into the output)
    rows = fichier("evenements")
    if rows is not None:
        import datetime as dt
        drop = {"quartier autre": 0, "titre vide": 0, "date absente ou passée": 0}
        n, pts = 0, spread(64, A_.tous, "evenements")
        for r in rows:
            if "saint roch" not in norm(r["Quartier"]) and "st roch" not in norm(r["Quartier"]):
                drop["quartier autre"] += 1
                continue
            titre = texte(r["Titre PROJET"])
            if not titre:
                drop["titre vide"] += 1
                continue
            d1, d2 = r["Date et heure de début"], r["Date et heure de fin"]
            if not isinstance(d1, dt.datetime) or d1 < dt.datetime.now():
                drop["date absente ou passée"] += 1
                continue
            if not isinstance(d2, dt.datetime):
                d2 = d1 + dt.timedelta(hours=2)
            site = texte(r["Site(s) de l'activité"])
            pt = A_.adresse(site) or pts[n % len(pts)]
            cat = texte(r["Catégorie"])
            th = ["culture"] if re.search(r"cultur|art|spectacle|musique", norm(cat)) else ["sport"] if "sport" in norm(cat) else ["communautaire"]
            evs.append(dict(id=f"fourni-{hid('ev', titre, d1.isoformat())}", titre=titre, lieuNom=site or "Saint-Roch", lat=pt[0], lng=pt[1],
                            debut=d1.strftime("%Y-%m-%dT%H:%M:00"), fin=d2.strftime("%Y-%m-%dT%H:%M:00"), categorie=cat or "événement",
                            themes=th, source="fourni", ouvertAuPublic=True))
            n += 1
        rapport.append(f"  -> {n} événements écrits; rejets: {drop}")

    out.mkdir(parents=True, exist_ok=True)
    for nom, data in (("profils-reels.json", profils), ("lieux-reels.json", lieux), ("evenements-reels.json", evs)):
        json.dump(data, open(out / nom, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("\n".join(rapport))
    print(f"écrit dans {out}: {len(profils)} profils, {len(lieux)} lieux, {len(evs)} événements (aucune valeur affichée ici)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
