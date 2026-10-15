"""Schema-only inventory of the organizers' data pack.

Prints file names, sizes, formats, sheet/layer names, row counts, column names
and inferred type words. It never prints a cell value, coordinate or PDF text.
Usage: py -3.12 scripts/schema.py "<folder>"
"""
import datetime as dt
import json
import os
import sys
import zipfile

SAMPLE = 500  # rows examined per column for type inference


def type_word(v):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return "empty"
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, int):
        return "int"
    if isinstance(v, float):
        return "int" if v.is_integer() else "float"
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return "date"
    return "text"


def merge(types):
    kinds = {t for t in types if t != "empty"}
    if not kinds:
        return "empty"
    if kinds == {"int", "float"}:
        return "float"
    return kinds.pop() if len(kinds) == 1 else "mixed(" + "/".join(sorted(kinds)) + ")"


def header_unclear(cells):
    filled = [c for c in cells if c is not None and str(c).strip() != ""]
    if len(filled) < 2:
        return True
    if any(not isinstance(c, str) for c in filled):
        return True
    if any(len(c) > 80 for c in filled):
        return True
    return False


def xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        first = next(rows, None)
        if first is None:
            print(f"  sheet '{ws.title}': empty sheet")
            continue
        first = list(first)
        while first and (first[-1] is None or str(first[-1]).strip() == ""):
            first.pop()
        n = 0
        cols = [[] for _ in first]
        for r in rows:
            if all(c is None or (isinstance(c, str) and not c.strip()) for c in r):
                continue
            n += 1
            if n <= SAMPLE:
                for i in range(len(first)):
                    cols[i].append(type_word(r[i] if i < len(r) else None))
        print(f"  sheet '{ws.title}': {n} data rows, {len(first)} columns")
        if header_unclear(first):
            print("    header row unclear")
            continue
        for i, name in enumerate(first):
            print(f"    - {name} [{merge(cols[i])}]")


def geojson(path):
    with open(path, encoding="utf-8") as f:
        d = json.load(f)
    feats = d.get("features", [])
    geoms = sorted({(x.get("geometry") or {}).get("type", "none") for x in feats})
    print(f"  features: {len(feats)}; geometry types: {', '.join(geoms)}")
    props = {}
    for x in feats[:SAMPLE]:
        for k, v in (x.get("properties") or {}).items():
            props.setdefault(k, []).append(type_word(v))
    for k, t in props.items():
        print(f"    - {k} [{merge(t)}]")


def gdb(vsi, label):
    import pyogrio
    for name, geom in pyogrio.list_layers(vsi):
        info = pyogrio.read_info(vsi, layer=name)
        print(f"  layer '{name}': features {info['features']}; geometry {geom}")
        for fld, dtp in zip(info["fields"], info["dtypes"]):
            print(f"    - {fld} [{dtp}]")


def zipf(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
    print(f"  inner entries: {len(names)}")
    gdbs = sorted({n.split("/")[0] for n in names if ".gdb" in n.split("/")[0]})
    for n in names:
        if len(names) <= 60 and not any(g in n for g in gdbs):
            print(f"    {n}")
    for g in gdbs:
        print(f"  geodatabase: {g}")
        gdb("/vsizip/" + path.replace("\\", "/") + "/" + g, g)
    shp = [n for n in names if n.lower().endswith((".shp", ".geojson", ".json", ".csv", ".xlsx"))]
    for n in shp:
        if n.lower().endswith(".shp"):
            import pyogrio
            vsi = "/vsizip/" + path.replace("\\", "/") + "/" + n
            info = pyogrio.read_info(vsi)
            print(f"  shapefile {n}: features {info['features']}; geometry {info['geometry_type']}")
            for fld, dtp in zip(info["fields"], info["dtypes"]):
                print(f"    - {fld} [{dtp}]")


def pdf(path):
    from pypdf import PdfReader
    print(f"  pages: {len(PdfReader(path).pages)}")


def main(folder):
    for name in sorted(os.listdir(folder)):
        p = os.path.join(folder, name)
        if not os.path.isfile(p):
            continue
        ext = os.path.splitext(name)[1].lower()
        print(f"\nFILE: {name} | {os.path.getsize(p)} bytes | {ext.lstrip('.')}")
        try:
            if ext == ".xlsx":
                xlsx(p)
            elif ext in (".geojson", ".json"):
                geojson(p)
            elif ext == ".zip":
                zipf(p)
            elif ext == ".pdf":
                pdf(p)
            else:
                print("  (format not handled)")
        except Exception as e:
            print(f"  ERROR {type(e).__name__}")


if __name__ == "__main__":
    main(sys.argv[1])
